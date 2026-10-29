"""File exports: a new-calendar ICS, a CSV, and an XLSX workbook with charts.

Each export is one file written by one CLI flag. The ICS and CSV need nothing
beyond the standard library; the workbook needs ``openpyxl`` (the ``[xlsx]``
extra) and says so plainly when it is missing.

Safety rules shared by all three:

* A path must name a regular file (or nothing yet) in an existing directory.
  A symlink is refused rather than followed, so an export cannot be steered
  into overwriting a file somewhere else.
* Content is built in memory first and written to a temporary file in the
  destination directory, then moved into place with ``os.replace``, so a
  failure never leaves a half-written file and never touches the old one.
* Text that a spreadsheet would run as a formula is neutralized (see
  :func:`neutralize`).
"""

from __future__ import annotations

import csv
import io
import os
import stat
import tempfile
from collections import OrderedDict
from typing import Any, Dict, List, Optional, Sequence

from .catalog import ResolvedHoliday
from .config import CompanyConfig
from .efficiency import ReturnCurve
from .planner import PtoPlan
from .report import GoalSummary, _named_days, render_ics

__all__ = [
    "ExportError",
    "CSV_HEADER",
    "neutralize",
    "render_calendar_ics",
    "render_csv",
    "write_exports",
]

XLSX_HINT = (
    "--export-xlsx needs the optional openpyxl package: "
    "pip install 'pto-colada[xlsx]'"
)

CSV_HEADER = ["kind", "start", "end", "length_days", "pto_cost", "efficiency", "description"]

#: First characters that make a spreadsheet read a cell as a formula (OWASP
#: CSV injection list, which adds tab and carriage return).
_FORMULA_LEADERS = ("=", "+", "-", "@", "\t", "\r")


class ExportError(ValueError):
    """An export could not be made; the message says why and what to do."""


def neutralize(value: Any) -> Any:
    """Prefix text that starts with a formula character so it stays text.

    Only strings are touched. The prefix is a single quote, the convention
    spreadsheets treat as "this is text".
    """
    if isinstance(value, str) and value.startswith(_FORMULA_LEADERS):
        return "'" + value
    return value


# --------------------------------------------------------------------------- ics


def render_calendar_ics(
    config: CompanyConfig,
    holidays_by_year: Dict[int, List[ResolvedHoliday]],
    plan: Optional[PtoPlan],
    include_holidays: bool = True,
    goal: Optional[GoalSummary] = None,
) -> str:
    """An ICS meant to be imported as its own calendar.

    It is the report's all-day-event calendar (holidays, PTO days, breaks) with
    its ``X-WR-CALNAME``, so importing creates a calendar named
    "<company> time off" and deleting that calendar removes every event.
    """
    return render_ics(
        config, holidays_by_year, plan,
        include_holidays=include_holidays, efficiency=True, goal=goal,
    )


# --------------------------------------------------------------------------- rows


def _require_plan(plan: Optional[PtoPlan], flag: str) -> PtoPlan:
    if plan is None:
        raise ExportError(f"{flag} needs a PTO plan: add --pto")
    return plan


def _ratio(value: Optional[float]) -> Any:
    return round(value, 3) if value is not None else ""


def plan_rows(plan: PtoPlan, config: CompanyConfig) -> List[List[Any]]:
    """Breaks first, then PTO days, as plain values (dates as ISO strings)."""
    named = _named_days(config)
    rows: List[List[Any]] = []
    for item in plan.breaks:
        rows.append([
            "break", item.start.isoformat(), item.end.isoformat(), item.length,
            item.pto_cost, _ratio(item.efficiency),
            "; ".join(OrderedDict.fromkeys(item.holidays)),
        ])
    for date in plan.pto_dates:
        rows.append([
            "pto", date.isoformat(), date.isoformat(), 1, 1, "", named.get(date, ""),
        ])
    return rows


def render_csv(plan: Optional[PtoPlan], config: CompanyConfig) -> str:
    """Breaks and PTO days as CSV (CRLF line ends, formula-safe text)."""
    plan = _require_plan(plan, "--export-csv")
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow(CSV_HEADER)
    for row in plan_rows(plan, config):
        writer.writerow([neutralize(cell) for cell in row])
    return out.getvalue()


# --------------------------------------------------------------------------- xlsx


def pto_per_month(plan: PtoPlan) -> List[Any]:
    """``(YYYY-MM, count)`` for every month the plan's window touches."""
    if plan.days:
        first = min(d.date for d in plan.days)
        last = max(d.date for d in plan.days)
    elif plan.pto_dates:
        first, last = min(plan.pto_dates), max(plan.pto_dates)
    else:
        return []
    counts: Dict[str, int] = {}
    year, month = first.year, first.month
    while (year, month) <= (last.year, last.month):
        counts["%04d-%02d" % (year, month)] = 0
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    for date in plan.pto_dates:
        key = "%04d-%02d" % (date.year, date.month)
        counts[key] = counts.get(key, 0) + 1
    return sorted(counts.items())


def _load_openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError:
        raise ExportError(XLSX_HINT) from None
    return openpyxl


def build_workbook(
    plan: PtoPlan,
    config: CompanyConfig,
    curve: Optional[ReturnCurve],
) -> Any:
    """The workbook: Summary, Breaks, PTO, PTO per month and Return curve sheets.

    Charts reference the sheet cells, so editing a number in the workbook moves
    the chart: a return curve (days in breaks per budget), a breaks timeline
    (each break's length and PTO cost) and PTO per month.
    """
    openpyxl = _load_openpyxl()
    from openpyxl.chart import BarChart, LineChart, Reference

    wb = openpyxl.Workbook()

    def put(ws: Any, row: Sequence[Any]) -> None:
        ws.append(list(row))
        for cell in ws[ws.max_row]:
            # Text is stored as text: openpyxl would otherwise turn any string
            # starting with "=" into a formula.
            if isinstance(cell.value, str):
                cell.data_type = "s"

    # Summary
    summary = wb.active
    summary.title = "Summary"
    put(summary, ["Item", "Value"])
    put(summary, ["Calendar", config.name])
    efficiency = plan.efficiency
    put(summary, ["Efficiency score (days in breaks per PTO day)",
                  round(efficiency, 3) if efficiency is not None else "n/a"])
    put(summary, ["PTO days booked (floating included)", plan.total_pto])
    put(summary, ["Days in breaks", plan.days_in_breaks])
    put(summary, ["Breaks", len(plan.breaks)])
    put(summary, ["Minimum break length", plan.min_break])
    knee = curve.knee if curve is not None else None
    put(summary, ["Return curve knee (PTO per year)", knee.budget if knee else "n/a"])
    put(summary, [])
    put(summary, ["Year", "Granted", "Spent", "Remaining"])
    for year in sorted(plan.spent_by_year):
        put(summary, [year, plan.granted_total(year), plan.spent_by_year[year], plan.remaining(year)])
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 24

    # Breaks (the timeline chart reads this sheet)
    breaks = wb.create_sheet("Breaks")
    put(breaks, ["Start", "End", "Days off", "PTO cost", "Efficiency", "Holidays"])
    for item in plan.breaks:
        put(breaks, [item.start.isoformat(), item.end.isoformat(), item.length,
                     item.pto_cost, _ratio(item.efficiency),
                     "; ".join(OrderedDict.fromkeys(item.holidays))])

    # PTO days
    pto = wb.create_sheet("PTO")
    put(pto, ["Date", "Reason"])
    named = _named_days(config)
    for date in plan.pto_dates:
        put(pto, [date.isoformat(), named.get(date, "")])

    # PTO per month
    monthly = wb.create_sheet("PTO per month")
    put(monthly, ["Month", "PTO days"])
    months = pto_per_month(plan)
    for label, count in months:
        put(monthly, [label, count])

    # Return curve
    curve_sheet = None
    if curve is not None:
        curve_sheet = wb.create_sheet("Return curve")
        put(curve_sheet, ["Budget per year", "PTO spent", "Days in breaks", "Marginal"])
        for p in curve.points:
            put(curve_sheet, [p.budget, p.pto_spent, p.days_in_breaks,
                              round(p.marginal, 3) if p.marginal is not None else ""])

    def style(chart: Any, title: str, x: str, y: str, anchor_to: Any, width: float = 18) -> None:
        chart.title = title
        chart.x_axis.title = x
        chart.y_axis.title = y
        chart.width, chart.height = width, 9
        # openpyxl 3.1 hides axes unless told otherwise.
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        anchor_to.add_chart(chart, "H2")

    if curve_sheet is not None and curve.points:
        line = LineChart()
        n = len(curve.points) + 1
        line.add_data(Reference(curve_sheet, min_col=3, min_row=1, max_row=n), titles_from_data=True)
        line.set_categories(Reference(curve_sheet, min_col=1, min_row=2, max_row=n))
        style(line, "Return curve: days in breaks by PTO budget",
              "PTO budget per year", "Days in breaks", curve_sheet)

    if plan.breaks:
        bars = BarChart()
        bars.type = "col"
        n = len(plan.breaks) + 1
        bars.add_data(Reference(breaks, min_col=3, max_col=4, min_row=1, max_row=n),
                      titles_from_data=True)
        bars.set_categories(Reference(breaks, min_col=1, min_row=2, max_row=n))
        style(bars, "Breaks timeline", "Break start", "Days", breaks, width=24)

    if months:
        per_month = BarChart()
        per_month.type = "col"
        n = len(months) + 1
        per_month.add_data(Reference(monthly, min_col=2, min_row=1, max_row=n),
                           titles_from_data=True)
        per_month.set_categories(Reference(monthly, min_col=1, min_row=2, max_row=n))
        style(per_month, "PTO per month", "Month", "PTO days", monthly, width=24)

    return wb


# --------------------------------------------------------------------------- writing


def check_target(path: str, flag: str) -> str:
    """Return the path to write, else raise ``ExportError``.

    The directory part is resolved (so a symlinked directory is followed to its
    real location); the final component must not be a symlink and, if it
    exists, must be a regular file.
    """
    if not path or "\x00" in path:
        raise ExportError(f"{flag}: not a usable file path")
    absolute = os.path.abspath(os.path.expanduser(path))
    parent = os.path.realpath(os.path.dirname(absolute))
    if not os.path.isdir(parent):
        raise ExportError(f"{flag}: directory does not exist: {os.path.dirname(absolute)}")
    target = os.path.join(parent, os.path.basename(absolute))
    if os.path.islink(target):
        raise ExportError(f"{flag}: {path} is a symbolic link; refusing to write through it")
    try:
        mode = os.lstat(target).st_mode
    except FileNotFoundError:
        return target
    except OSError as exc:
        raise ExportError(f"{flag}: cannot inspect {path}: {exc.strerror or exc}") from None
    if not stat.S_ISREG(mode):
        raise ExportError(f"{flag}: {path} exists and is not a regular file")
    return target


def _same_file(a: str, b: str) -> bool:
    if a == b:
        return True
    try:
        return os.path.samefile(a, b)  # hard links and the like; False if either is missing
    except OSError:
        return False


def _atomic_write(target: str, writer: Any, flag: str) -> None:
    """Run ``writer(tmp_path)`` then rename the result over ``target``.

    The temporary file is created exclusively in the destination directory
    (so the rename is atomic) and removed on any failure. ``os.replace``
    swaps the directory entry itself, so even a symlink planted after the
    check is replaced rather than followed. An existing file keeps its
    permissions; a new one gets the usual umask-filtered 0666.
    """
    tmp = None
    try:
        fd, tmp = tempfile.mkstemp(
            prefix=".lwp-export-", suffix=".tmp", dir=os.path.dirname(target)
        )
        os.close(fd)
        writer(tmp)
        try:
            mode = stat.S_IMODE(os.lstat(target).st_mode)
        except FileNotFoundError:
            umask = os.umask(0)
            os.umask(umask)
            mode = 0o666 & ~umask
        os.chmod(tmp, mode)
        os.replace(tmp, target)
    except OSError as exc:
        raise ExportError(f"{flag}: could not write {target}: {exc.strerror or exc}") from None
    finally:
        if tmp is not None and os.path.exists(tmp):
            os.unlink(tmp)


def _write_bytes(data: bytes):
    def write(tmp: str) -> None:
        with open(tmp, "wb") as handle:
            handle.write(data)
    return write


def write_exports(
    config: CompanyConfig,
    holidays_by_year: Dict[int, List[ResolvedHoliday]],
    plan: Optional[PtoPlan],
    goal: Optional[GoalSummary] = None,
    curve: Optional[ReturnCurve] = None,
    include_holidays: bool = True,
    ics: Optional[str] = None,
    csv_path: Optional[str] = None,
    xlsx: Optional[str] = None,
) -> List[str]:
    """Write every requested export; return the paths written.

    Everything is checked and built before the first file is written, so a bad
    path, a missing dependency or a missing plan leaves nothing behind.
    """
    jobs = []
    requested = [(ics, "--export-ics"), (csv_path, "--export-csv"), (xlsx, "--export-xlsx")]
    seen: Dict[str, str] = {}
    for path, flag in requested:
        if path is None:
            continue
        target = check_target(path, flag)
        for other, other_flag in seen.items():
            if _same_file(target, other):
                raise ExportError(f"{flag} and {other_flag} name the same file: {path}")
        seen[target] = flag
        jobs.append((flag, target))
    targets = {flag: target for flag, target in jobs}

    payloads = []
    if "--export-ics" in targets:
        text = render_calendar_ics(config, holidays_by_year, plan, include_holidays, goal)
        # newline="" semantics: the text already carries CRLF line ends.
        payloads.append((targets["--export-ics"], "--export-ics", _write_bytes(text.encode("utf-8"))))
    if "--export-csv" in targets:
        text = render_csv(plan, config)
        payloads.append((targets["--export-csv"], "--export-csv", _write_bytes(text.encode("utf-8"))))
    if "--export-xlsx" in targets:
        workbook = build_workbook(_require_plan(plan, "--export-xlsx"), config, curve)
        payloads.append((targets["--export-xlsx"], "--export-xlsx", workbook.save))

    for target, flag, writer in payloads:
        _atomic_write(target, writer, flag)
    return [target for target, _flag, _writer in payloads]
